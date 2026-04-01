from rest_framework.views import APIView
from rest_framework.response import Response
from rest_framework.permissions import IsAuthenticated
from rest_framework.parsers import MultiPartParser
from rest_framework import status
from django.db.models import Count, Q, Sum
from django.http import HttpResponse
from .etl import cargar_excel, cargar_excel_facturacion, limpiar_datos, limpiar_facturacion_externa, ejecutar_recategorizacion, _safe_float, _safe_date
from .models import ResultadoImportacion, Cliente, ResumenTarifario, ClienteNoApto, Anomalia, ConsumoMensual, MatrizRecategorizacion
import pandas as pd
import io


class ImportarExcelView(APIView):
    permission_classes = [IsAuthenticated]
    parser_classes = [MultiPartParser]

    def post(self, request):
        archivos_lectura     = request.FILES.getlist("archivos_lectura")
        archivos_facturacion = request.FILES.getlist("archivos_facturacion")

        if not archivos_lectura:
            return Response({"error": "Se requiere al menos un archivo de lecturas"}, status=status.HTTP_400_BAD_REQUEST)
        if not archivos_facturacion:
            return Response({"error": "Se requiere al menos un archivo de facturación"}, status=status.HTTP_400_BAD_REQUEST)

        try:
            dfs_lectura     = []
            dfs_facturacion = []
            for archivo in archivos_lectura:
                df_l, df_f = cargar_excel(archivo)
                dfs_lectura.append(df_l)
                dfs_facturacion.append(df_f)

            df_lectura     = pd.concat(dfs_lectura,     ignore_index=True)
            df_facturacion = pd.concat(dfs_facturacion, ignore_index=True)
            df_lectura, df_facturacion = limpiar_datos(df_lectura, df_facturacion)

            df_facturacion_externa = None
            if archivos_facturacion:
                dfs_fact_ext = []
                for archivo in archivos_facturacion:
                    df_fe = cargar_excel_facturacion(archivo)
                    dfs_fact_ext.append(df_fe)
                df_facturacion_externa = limpiar_facturacion_externa(
                    pd.concat(dfs_fact_ext, ignore_index=True)
                )

            cuadro_2, cuadro_3, cuadro_4, cuadro_5, df_mensual, cuadro_matriz = ejecutar_recategorizacion(
                df_lectura, df_facturacion, df_facturacion_externa
            )

            total_clientes  = len(cuadro_2)
            recategorizados = int((cuadro_2["Estado"] == "Recategorizado").sum())
            sin_cambios     = int((cuadro_2["Estado"] == "Sin cambio").sum())
            no_aptos        = cuadro_4["Instalación"].nunique()
            anomalias       = len(cuadro_5)

            importacion = ResultadoImportacion.objects.create(
                usuario         = request.user,
                total_registros = total_clientes + no_aptos,
                procesados      = total_clientes,
                recategorizados = recategorizados,
                sin_cambios     = sin_cambios,
                no_aptos        = no_aptos,
                anomalias       = anomalias,
            )

            clientes_bulk = []
            for _, row in cuadro_2.iterrows():
                clientes_bulk.append(Cliente(
                    importacion         = importacion,
                    instalacion         = str(row["Instalación"]),
                    cuenta_contrato     = str(row["Cuenta_contrato"]),
                    total_dias_consumo  = float(row["Total_dias_consumo"] or 0),
                    total_consumo       = float(row["Total_consumo_facturado"] or 0),
                    promedio_diario     = float(row["Promedio_diario"] or 0),
                    promedio_mensual    = float(row["Promedio_mensual_redondeado"] or 0),
                    tarifa_anterior     = str(row["Tarifa_referencia"] or ""),
                    tarifa_nueva        = str(row["Nueva_tarifa"] or ""),
                    estado              = str(row["Estado"] or ""),
                    porcion             = str(row["Porcion"]) if pd.notna(row["Porcion"]) else None,
                    unidad_predial      = str(row["Unidad_Predial"]) if pd.notna(row["Unidad_Predial"]) else None,
                    rango_consumo       = str(row["Rango_consumo"]) if pd.notna(row.get("Rango_consumo")) else None,
                ))
            Cliente.objects.bulk_create(clientes_bulk, batch_size=500)

            resumenes_bulk = []
            for _, row in cuadro_3.iterrows():
                resumenes_bulk.append(ResumenTarifario(
                    importacion       = importacion,
                    tarifa_anterior   = str(row["Tarifa_referencia"]),
                    tarifa_nueva      = str(row["Nueva_tarifa"]),
                    cantidad_clientes = int(row["Cantidad_clientes"]),
                    porcentaje        = float(row["Porcentaje"]),
                ))
            ResumenTarifario.objects.bulk_create(resumenes_bulk)

            no_aptos_bulk = []
            for _, row in cuadro_4.iterrows():
                no_aptos_bulk.append(ClienteNoApto(
                    importacion       = importacion,
                    cuenta_contrato   = str(row["Cuenta contrato"] or ""),
                    instalacion       = str(row["Instalación"] or ""),
                    tarifa_referencia = str(row["Tarifa referencia"]) if pd.notna(row.get("Tarifa referencia")) else None,
                    observacion       = str(row["Observación"] or ""),
                    porcion           = str(row["Porción"]) if pd.notna(row.get("Porción")) else None,
                    unidad_predial    = str(row["Unidad Predial"]) if pd.notna(row.get("Unidad Predial")) else None,
                    meses_en_ventana  = float(row["meses_en_ventana"]) if pd.notna(row.get("meses_en_ventana")) else None,
                    estado_inicial    = str(row["Estado inicial"]) if pd.notna(row.get("Estado inicial")) else None,
                ))
            ClienteNoApto.objects.bulk_create(no_aptos_bulk, batch_size=500)

            if len(cuadro_5) > 0:
                anomalias_bulk = []
                for _, row in cuadro_5.iterrows():
                    anomalias_bulk.append(Anomalia(
                        importacion     = importacion,
                        cuenta_contrato = str(row["Cuenta_contrato"] or ""),
                        instalacion     = str(row["Instalacion"] or ""),
                        fecha           = row["Fecha"] if pd.notna(row["Fecha"]) else None,
                        tipo_anomalia   = str(row["Tipo_anomalia"] or ""),
                    ))
                Anomalia.objects.bulk_create(anomalias_bulk, batch_size=500)

            clientes_map = {
                c.cuenta_contrato: c
                for c in Cliente.objects.filter(importacion=importacion)
            }

            consumo_bulk = []
            for _, row in df_mensual.iterrows():
                cuenta = str(row["Cuenta_contrato"])
                consumo_bulk.append(ConsumoMensual(
                    importacion     = importacion,
                    cliente         = clientes_map.get(cuenta),
                    cuenta_contrato = cuenta,
                    instalacion     = str(row["Instalación"]),
                    porcion         = str(row["Porcion"]) if pd.notna(row.get("Porcion")) else None,
                    periodo         = str(row["Periodo"]),
                    consumo         = float(row["Consumo_mes"] or 0),
                    dias            = int(row["Dias_mes"] or 0),
                    tarifa          = str(row["Tarifa_referencia"]) if pd.notna(row.get("Tarifa_referencia")) else None,
                ))
            ConsumoMensual.objects.bulk_create(consumo_bulk, batch_size=1000)

            matriz_bulk = []
            for _, row in cuadro_matriz.iterrows():
                def sf(col): return _safe_float(row.get(col))
                def sd(col): return _safe_date(row.get(col))
                matriz_bulk.append(MatrizRecategorizacion(
                    importacion         = importacion,
                    cuenta_contrato     = str(row.get("Cuenta_contrato") or ""),
                    instalacion         = str(row.get("Instalación") or ""),
                    porcion             = str(row["Porcon"]) if pd.notna(row.get("Porcon")) else None,
                    tipo_tarifa         = str(row.get("Tarifa_referencia") or ""),
                    cf_mes_historico    = sf("CF_mes_historico"),
                    fl_mes_historico    = sd("FL_mes_historico"),
                    cl_mes_historico    = sf("CL_mes_historico"),
                    cf_mes_1            = sf("CF_mes_1"),
                    fl_mes_1            = sd("FL_mes_1"),
                    dc_mes_1            = sf("DC_mes_1"),
                    cl_mes_1            = sf("CL_mes_1"),
                    tarifa_mes_1        = str(row["Tarifa_mes_1"]) if pd.notna(row.get("Tarifa_mes_1")) else None,
                    cf_mes_2            = sf("CF_mes_2"),
                    fl_mes_2            = sd("FL_mes_2"),
                    dc_mes_2            = sf("DC_mes_2"),
                    cl_mes_2            = sf("CL_mes_2"),
                    tarifa_mes_2        = str(row["Tarifa_mes_2"]) if pd.notna(row.get("Tarifa_mes_2")) else None,
                    cf_mes_3            = sf("CF_mes_3"),
                    fl_mes_3            = sd("FL_mes_3"),
                    dc_mes_3            = sf("DC_mes_3"),
                    cl_mes_3            = sf("CL_mes_3"),
                    tarifa_mes_3        = str(row["Tarifa_mes_3"]) if pd.notna(row.get("Tarifa_mes_3")) else None,
                    cf_mes_4            = sf("CF_mes_4"),
                    fl_mes_4            = sd("FL_mes_4"),
                    dc_mes_4            = sf("DC_mes_4"),
                    cl_mes_4            = sf("CL_mes_4"),
                    tarifa_mes_4        = str(row["Tarifa_mes_4"]) if pd.notna(row.get("Tarifa_mes_4")) else None,
                    cf_mes_5            = sf("CF_mes_5"),
                    fl_mes_5            = sd("FL_mes_5"),
                    dc_mes_5            = sf("DC_mes_5"),
                    cl_mes_5            = sf("CL_mes_5"),
                    tarifa_mes_5        = str(row["Tarifa_mes_5"]) if pd.notna(row.get("Tarifa_mes_5")) else None,
                    cf_mes_6            = sf("CF_mes_6"),
                    fl_mes_6            = sd("FL_mes_6"),
                    dc_mes_6            = sf("DC_mes_6"),
                    cl_mes_6            = sf("CL_mes_6"),
                    lectura_anterior    = sf("Lectura Anterior"),
                    lectura_actual      = sf("Lectura Actual"),
                    factor_correccion   = sf("Factor de Corrección"),
                    total_consumo       = float(row.get("Total_consumo_facturado") or 0),
                    total_dias          = float(row.get("Total_dias_consumo") or 0),
                    promedio_diario     = float(row.get("Promedio_diario") or 0),
                    promedio_mensual    = float(row.get("Promedio_mensual") or 0),
                    promedio_redondeado = int(row.get("Promedio_redondeado") or 0),
                    tarifa_actual       = str(row.get("Tarifa_referencia") or ""),
                    tarifa_nueva        = str(row.get("Tarifa_nueva") or ""),
                    recategorizar       = str(row.get("Recategorizar") or "No"),
                    rango_consumo       = str(row.get("Rango_consumo") or ""),
                    comportamiento      = str(row.get("Comportamiento") or ""),
                ))
            MatrizRecategorizacion.objects.bulk_create(matriz_bulk, batch_size=500)

            return Response({
                "total":                   total_clientes + no_aptos,
                "procesados":              total_clientes,
                "recategorizados":         recategorizados,
                "sin_cambios":             sin_cambios,
                "no_aptos":                no_aptos,
                "errores":                 0,
                "anomalias":               anomalias,
                "uso_facturacion_externa": bool(archivos_facturacion),
                "distribucion_categorias": cuadro_2["Nueva_tarifa"].value_counts().to_dict(),
                "cambios_tarifarios":      cuadro_3.to_dict("records"),
            }, status=status.HTTP_200_OK)

        except Exception as e:
            import traceback
            error_detalle = traceback.format_exc()
            print(error_detalle)
            return Response(
                {"error": f"Error procesando archivos: {str(e)}", "detalle": error_detalle},
                status=status.HTTP_500_INTERNAL_SERVER_ERROR
            )


class DashboardStatsView(APIView):
    permission_classes = [IsAuthenticated]

    def get(self, request):
        es_admin       = request.user.rol == 'admin'
        importacion_id = request.query_params.get('importacion_id', '')

        if es_admin:
            if importacion_id:
                try:
                    ultima = ResultadoImportacion.objects.get(id=importacion_id)
                except ResultadoImportacion.DoesNotExist:
                    ultima = ResultadoImportacion.objects.order_by('-fecha_importacion').first()
            else:
                ultima = ResultadoImportacion.objects.order_by('-fecha_importacion').first()
        else:
            if importacion_id:
                try:
                    ultima = ResultadoImportacion.objects.get(id=importacion_id, usuario=request.user)
                except ResultadoImportacion.DoesNotExist:
                    ultima = ResultadoImportacion.objects.filter(usuario=request.user).order_by('-fecha_importacion').first()
            else:
                ultima = ResultadoImportacion.objects.filter(usuario=request.user).order_by('-fecha_importacion').first()

        if not ultima:
            return Response({
                "es_admin":                es_admin,
                "total_clientes":          0,
                "recategorizados":         0,
                "sin_cambios":             0,
                "no_aptos":                0,
                "anomalias":               0,
                "distribucion_categorias": [],
                "cambios_tarifarios":      [],
                "no_aptos_observaciones":  [],
                "anomalias_por_tipo":      [],
                "consumo_por_periodo":     [],
                "distribucion_rangos":     [],
                "mis_importaciones":       [],
                "total_importaciones":     0,
                "usuarios_activos":        0,
            })

        dist = (
            Cliente.objects.filter(importacion=ultima)
            .values("tarifa_nueva")
            .annotate(cantidad=Count("id"))
        )

        cambios = (
            ResumenTarifario.objects.filter(importacion=ultima)
            .values("tarifa_anterior", "tarifa_nueva", "cantidad_clientes", "porcentaje")
        )

        no_aptos_obs = (
            ClienteNoApto.objects.filter(importacion=ultima)
            .values("observacion")
            .annotate(cantidad=Count("id"))
            .order_by("-cantidad")[:6]
        )

        anomalias_tipo = (
            Anomalia.objects.filter(importacion=ultima)
            .values("tipo_anomalia")
            .annotate(cantidad=Count("id"))
            .order_by("-cantidad")[:6]
        )

        consumo_por_periodo = list(
            ConsumoMensual.objects.filter(importacion=ultima)
            .values("periodo")
            .annotate(consumo_total=Sum("consumo"), dias_total=Sum("dias"))
            .order_by("periodo")
        )

        def _rango(promedio):
            if promedio <= 30:  return "REG-A1-CO (0-30 m³)"
            if promedio <= 300: return "REG-A2-CO (31-300 m³)"
            return "REG-B-CO (>300 m³)"

        clientes_qs = Cliente.objects.filter(importacion=ultima).values("promedio_mensual")
        rangos = {}
        for c in clientes_qs:
            rango = _rango(c["promedio_mensual"] or 0)
            rangos[rango] = rangos.get(rango, 0) + 1

        distribucion_rangos = [{"rango": k, "cantidad": v} for k, v in sorted(rangos.items())]

        data = {
            "es_admin":                es_admin,
            "importacion_id":          ultima.id,
            "total_clientes":          ultima.total_registros,
            "recategorizados":         ultima.recategorizados,
            "sin_cambios":             ultima.sin_cambios,
            "no_aptos":                ultima.no_aptos,
            "anomalias":               ultima.anomalias,
            "distribucion_categorias": list(dist),
            "cambios_tarifarios":      list(cambios),
            "no_aptos_observaciones":  list(no_aptos_obs),
            "anomalias_por_tipo":      list(anomalias_tipo),
            "consumo_por_periodo":     consumo_por_periodo,
            "distribucion_rangos":     distribucion_rangos,
        }

        if not es_admin:
            from zoneinfo import ZoneInfo
            mis_importaciones = []
            for imp in ResultadoImportacion.objects.filter(usuario=request.user).order_by('-fecha_importacion'):
                mis_importaciones.append({
                    "id"              : imp.id,
                    "fecha"           : imp.fecha_importacion.astimezone(ZoneInfo('America/Lima')).strftime("%d/%m/%Y %H:%M"),
                    "total_registros" : imp.total_registros,
                    "recategorizados" : imp.recategorizados,
                })
            data["mis_importaciones"] = mis_importaciones

        if es_admin:
            from django.contrib.auth import get_user_model
            from zoneinfo import ZoneInfo
            User = get_user_model()
            totales = ResultadoImportacion.objects.aggregate(
                total_registros       = Sum('total_registros'),
                total_recategorizados = Sum('recategorizados'),
            )
            data["globales"] = {
                "total_importaciones"  : ResultadoImportacion.objects.count(),
                "total_registros"      : totales["total_registros"] or 0,
                "total_recategorizados": totales["total_recategorizados"] or 0,
            }
            data["total_importaciones"] = data["globales"]["total_importaciones"]
            data["usuarios_activos"]    = User.objects.filter(activo=True).count()

            usuarios_lista = []
            for u in User.objects.filter(activo=True, rol='usuario').order_by('nombre'):
                importaciones_u = []
                for imp in ResultadoImportacion.objects.filter(usuario=u).order_by('-fecha_importacion'):
                    importaciones_u.append({
                        "id"             : imp.id,
                        "fecha_str"      : imp.fecha_importacion.astimezone(ZoneInfo('America/Lima')).strftime("%d/%m/%Y %H:%M"),
                        "total_registros": imp.total_registros,
                        "recategorizados": imp.recategorizados,
                    })
                if importaciones_u:
                    usuarios_lista.append({
                        "id"                 : u.id,
                        "nombre"             : u.nombre,
                        "apellido"           : u.apellido,
                        "total_importaciones": len(importaciones_u),
                        "importaciones"      : importaciones_u,
                    })
            data["usuarios_lista"] = usuarios_lista

            if ultima.usuario:
                data["importador"]        = f"{ultima.usuario.nombre} {ultima.usuario.apellido}"
                data["importacion_fecha"] = ultima.fecha_importacion.astimezone(ZoneInfo('America/Lima')).strftime("%d/%m/%Y %H:%M")
            else:
                data["importador"] = "Usuario eliminado"

        return Response(data)


class ClientesListView(APIView):
    permission_classes = [IsAuthenticated]

    def get(self, request):
        importacion_id = request.query_params.get('importacion_id', '')

        if importacion_id:
            try:
                ultima = ResultadoImportacion.objects.get(id=importacion_id, usuario=request.user)
            except ResultadoImportacion.DoesNotExist:
                ultima = ResultadoImportacion.objects.filter(usuario=request.user).order_by('-fecha_importacion').first()
        else:
            ultima = ResultadoImportacion.objects.filter(usuario=request.user).order_by('-fecha_importacion').first()

        if not ultima:
            return Response({"count": 0, "results": [], "porciones": [], "periodos": []})

        search      = request.query_params.get("search", "")
        porcion     = request.query_params.get("porcion", "")
        fecha_desde = request.query_params.get("fecha_desde", "")
        fecha_hasta = request.query_params.get("fecha_hasta", "")
        page        = int(request.query_params.get("page", 1))
        size        = int(request.query_params.get("page_size", 10))

        porciones_disponibles = list(
            ConsumoMensual.objects.filter(importacion=ultima)
            .exclude(porcion__isnull=True).exclude(porcion='')
            .values_list('porcion', flat=True)
            .distinct().order_by('porcion')
        )
        periodos_disponibles = list(
            ConsumoMensual.objects.filter(importacion=ultima)
            .values_list('periodo', flat=True)
            .distinct().order_by('periodo')
        )

        qs = Cliente.objects.filter(importacion=ultima)

        if search:
            qs = qs.filter(
                Q(instalacion__icontains=search) |
                Q(cuenta_contrato__icontains=search) |
                Q(porcion__icontains=search)
            )
        if porcion:
            qs = qs.filter(porcion=porcion)
        if fecha_desde or fecha_hasta:
            cm_qs = ConsumoMensual.objects.filter(importacion=ultima)
            if fecha_desde:
                cm_qs = cm_qs.filter(periodo__gte=fecha_desde)
            if fecha_hasta:
                cm_qs = cm_qs.filter(periodo__lte=fecha_hasta)
            cuentas_en_rango = cm_qs.values_list('cuenta_contrato', flat=True).distinct()
            qs = qs.filter(cuenta_contrato__in=cuentas_en_rango)

        total    = qs.count()
        offset   = (page - 1) * size
        clientes = qs[offset:offset + size]

        results = [{
            "id":                          c.id,
            "Instalación":                 c.instalacion,
            "Cuenta_contrato":             c.cuenta_contrato,
            "Total_dias_consumo":          round(c.total_dias_consumo, 2),
            "Total_consumo_facturado":     round(c.total_consumo, 2),
            "Tarifa_referencia":           c.tarifa_anterior,
            "Porcion":                     c.porcion or "-",
            "Unidad_Predial":              c.unidad_predial or "-",
            "Promedio_diario":             round(c.promedio_diario, 4),
            "Promedio_mensual":            round(c.promedio_mensual, 2),
            "Promedio_mensual_redondeado": int(c.promedio_mensual),
            "Nueva_tarifa":                c.tarifa_nueva,
            "Estado":                      c.estado,
        } for c in clientes]

        return Response({"count": total, "results": results, "porciones": porciones_disponibles, "periodos": periodos_disponibles})


class HistorialImportacionesView(APIView):
    permission_classes = [IsAuthenticated]

    def get(self, request):
        es_admin = request.user.rol == 'admin'

        if es_admin:
            importaciones = ResultadoImportacion.objects.order_by('-fecha_importacion')
        else:
            importaciones = ResultadoImportacion.objects.filter(usuario=request.user)

        data = []
        for imp in importaciones:
            item = {
                "id":              imp.id,
                "fecha":           imp.fecha_importacion.astimezone(
                    __import__('zoneinfo').ZoneInfo('America/Lima')
                ).strftime("%d/%m/%Y %H:%M"),
                "total_registros": imp.total_registros,
                "procesados":      imp.procesados,
                "recategorizados": imp.recategorizados,
                "sin_cambios":     imp.sin_cambios,
                "no_aptos":        imp.no_aptos,
                "anomalias":       imp.anomalias,
            }
            if es_admin:
                if imp.usuario:
                    item["usuario"] = f"{imp.usuario.nombre} {imp.usuario.apellido}"
                    item["email"]   = imp.usuario.email
                else:
                    item["usuario"] = "Usuario eliminado"
                    item["email"]   = "-"
            data.append(item)

        return Response(data)


class ExportarExcelView(APIView):
    permission_classes = [IsAuthenticated]

    def get(self, request):
        from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
        from openpyxl.utils import get_column_letter

        importacion_id = request.query_params.get("importacion_id")

        if importacion_id:
            if request.user.rol == 'admin':
                try:
                    imp = ResultadoImportacion.objects.get(id=importacion_id)
                except ResultadoImportacion.DoesNotExist:
                    return Response({"error": "Importación no encontrada"}, status=404)
            else:
                try:
                    imp = ResultadoImportacion.objects.get(id=importacion_id, usuario=request.user)
                except ResultadoImportacion.DoesNotExist:
                    return Response({"error": "Importación no encontrada"}, status=404)
        else:
            imp = ResultadoImportacion.objects.filter(usuario=request.user).order_by('-fecha_importacion').first()
            if not imp:
                return Response({"error": "No hay importaciones"}, status=404)

        from zoneinfo import ZoneInfo
        fecha_lima = imp.fecha_importacion.astimezone(ZoneInfo('America/Lima'))

        clientes  = list(Cliente.objects.filter(importacion=imp).values(
            "instalacion", "cuenta_contrato", "total_dias_consumo", "total_consumo",
            "promedio_diario", "promedio_mensual", "tarifa_anterior", "tarifa_nueva",
            "estado", "porcion", "unidad_predial"
        ))
        resumenes = list(ResumenTarifario.objects.filter(importacion=imp).values(
            "tarifa_anterior", "tarifa_nueva", "cantidad_clientes", "porcentaje"
        ))
        no_aptos  = list(ClienteNoApto.objects.filter(importacion=imp).values(
            "cuenta_contrato", "instalacion", "tarifa_referencia",
            "observacion", "porcion", "unidad_predial", "meses_en_ventana", "estado_inicial"
        ))
        anomalias = list(Anomalia.objects.filter(importacion=imp).values(
            "cuenta_contrato", "instalacion", "fecha", "tipo_anomalia"
        ))
        matriz = list(MatrizRecategorizacion.objects.filter(importacion=imp).values(
            "instalacion", "cuenta_contrato", "porcion", "tipo_tarifa",
            "cf_mes_historico", "fl_mes_historico", "cl_mes_historico",
            "cf_mes_1", "fl_mes_1", "dc_mes_1", "cl_mes_1", "tarifa_mes_1",
            "cf_mes_2", "fl_mes_2", "dc_mes_2", "cl_mes_2", "tarifa_mes_2",
            "cf_mes_3", "fl_mes_3", "dc_mes_3", "cl_mes_3", "tarifa_mes_3",
            "cf_mes_4", "fl_mes_4", "dc_mes_4", "cl_mes_4", "tarifa_mes_4",
            "cf_mes_5", "fl_mes_5", "dc_mes_5", "cl_mes_5", "tarifa_mes_5",
            "cf_mes_6", "fl_mes_6", "dc_mes_6", "cl_mes_6",
            "lectura_anterior", "lectura_actual", "factor_correccion",
            "total_consumo", "total_dias", "promedio_diario", "promedio_mensual",
            "promedio_redondeado", "tarifa_actual", "tarifa_nueva",
            "recategorizar", "rango_consumo", "comportamiento",
        ))

        # ── Colores ──
        AZUL_OSCURO = "0B1120"
        AZUL_MEDIO  = "1E3A5F"
        AMBAR       = "F59E0B"
        BLANCO      = "FFFFFF"
        GRIS_CLARO  = "F5F7FA"
        GRIS_MEDIO  = "E5E7EB"
        VERDE       = "10B981"
        VERDE_CLARO = "ECFDF5"
        ROJO_CLARO  = "FEF2F2"
        AMBAR_CLARO = "FFFBEB"
        AZUL_ROW    = "EFF6FF"

        thin  = Side(style="thin",   color=GRIS_MEDIO)
        thick = Side(style="medium", color=AZUL_MEDIO)
        border_thin   = Border(left=thin,  right=thin,  top=thin,  bottom=thin)
        border_header = Border(left=thick, right=thick, top=thick, bottom=thick)

        def estilo_header(ws, fila, col_inicio, col_fin, color_fondo=AZUL_OSCURO, color_letra=BLANCO):
            for col in range(col_inicio, col_fin + 1):
                cell = ws.cell(row=fila, column=col)
                cell.fill      = PatternFill("solid", fgColor=color_fondo)
                cell.font      = Font(bold=True, color=color_letra, size=10, name="Calibri")
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                cell.border    = border_header

        def titulo_hoja(ws, titulo, subtitulo, num_cols):
            ws.row_dimensions[1].height = 36
            ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=num_cols)
            c = ws.cell(row=1, column=1)
            c.value     = titulo
            c.fill      = PatternFill("solid", fgColor=AZUL_OSCURO)
            c.font      = Font(bold=True, color=AMBAR, size=14, name="Calibri")
            c.alignment = Alignment(horizontal="center", vertical="center")
            ws.row_dimensions[2].height = 22
            ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=num_cols)
            c2 = ws.cell(row=2, column=1)
            c2.value     = subtitulo
            c2.fill      = PatternFill("solid", fgColor=AZUL_MEDIO)
            c2.font      = Font(color=BLANCO, size=9, italic=True, name="Calibri")
            c2.alignment = Alignment(horizontal="center", vertical="center")
            ws.row_dimensions[3].height = 4
            ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=num_cols)
            ws.cell(row=3, column=1).fill = PatternFill("solid", fgColor=AMBAR)

        def autofit(ws, min_width=10, max_width=40):
            for col in ws.columns:
                max_len    = 0
                col_letter = get_column_letter(col[0].column)
                for cell in col:
                    try:
                        if cell.value:
                            max_len = max(max_len, len(str(cell.value)))
                    except:
                        pass
                ws.column_dimensions[col_letter].width = min(max(max_len + 2, min_width), max_width)

        output = io.BytesIO()
        with pd.ExcelWriter(output, engine="openpyxl") as writer:

            # ── Hoja 1: Clientes Recategorizados ──
            df1 = pd.DataFrame(clientes).rename(columns={
                "instalacion":        "Instalación",
                "cuenta_contrato":    "Cuenta Contrato",
                "total_dias_consumo": "Total Días",
                "total_consumo":      "Total Consumo m3",
                "promedio_diario":    "Prom. Diario",
                "promedio_mensual":   "Prom. Mensual",
                "tarifa_anterior":    "Tarifa Anterior",
                "tarifa_nueva":       "Tarifa Nueva",
                "estado":             "Estado",
                "porcion":            "Porción",
                "unidad_predial":     "Unidad Predial",
            })
            df1.to_excel(writer, sheet_name="Clientes Recategorizados", index=False, startrow=4)
            ws1 = writer.sheets["Clientes Recategorizados"]
            nc1 = len(df1.columns)
            titulo_hoja(ws1, "CONTUGAS — Clientes Recategorizados",
                        f"Importación: {fecha_lima.strftime('%d/%m/%Y %H:%M')}  |  Total: {len(df1):,} clientes", nc1)
            ws1.row_dimensions[5].height = 32
            estilo_header(ws1, 5, 1, nc1, AZUL_MEDIO, BLANCO)
            for i, row in enumerate(ws1.iter_rows(min_row=6, max_row=5 + len(df1), min_col=1, max_col=nc1)):
                zebra = i % 2 == 0
                for cell in row:
                    cell.fill      = PatternFill("solid", fgColor=AZUL_ROW if zebra else BLANCO)
                    cell.font      = Font(size=9, name="Calibri", color="374151")
                    cell.alignment = Alignment(vertical="center")
                    cell.border    = border_thin
                    if cell.column == df1.columns.get_loc("Estado") + 1:
                        if cell.value == "Recategorizado":
                            cell.fill = PatternFill("solid", fgColor=VERDE_CLARO)
                            cell.font = Font(size=9, name="Calibri", color=VERDE, bold=True)
                        elif cell.value == "Sin cambio":
                            cell.fill = PatternFill("solid", fgColor=GRIS_CLARO)
                            cell.font = Font(size=9, name="Calibri", color="6B7280", bold=True)
            ws1.freeze_panes = "A6"
            autofit(ws1)

            # ── Hoja 2: Resumen Tarifario ──
            df2 = pd.DataFrame(resumenes).rename(columns={
                "tarifa_anterior":   "Tarifa Anterior",
                "tarifa_nueva":      "Tarifa Nueva",
                "cantidad_clientes": "Cantidad Clientes",
                "porcentaje":        "Porcentaje %",
            })
            df2.to_excel(writer, sheet_name="Resumen Tarifario", index=False, startrow=4)
            ws2 = writer.sheets["Resumen Tarifario"]
            nc2 = len(df2.columns)
            titulo_hoja(ws2, "CONTUGAS — Resumen Tarifario",
                        f"Movimientos entre categorías  |  {fecha_lima.strftime('%d/%m/%Y %H:%M')}", nc2)
            ws2.row_dimensions[5].height = 32
            estilo_header(ws2, 5, 1, nc2, AZUL_OSCURO, AMBAR)
            for i, row in enumerate(ws2.iter_rows(min_row=6, max_row=5 + len(df2), min_col=1, max_col=nc2)):
                zebra = i % 2 == 0
                for cell in row:
                    cell.fill      = PatternFill("solid", fgColor=AZUL_ROW if zebra else BLANCO)
                    cell.font      = Font(size=9, name="Calibri", color="374151")
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                    cell.border    = border_thin
                    if cell.column == 3:
                        cell.font = Font(size=10, bold=True, name="Calibri", color=AZUL_MEDIO)
                tarifa_ant = ws2.cell(row=row[0].row, column=1).value
                tarifa_nva = ws2.cell(row=row[0].row, column=2).value
                if tarifa_ant == tarifa_nva:
                    for cell in row:
                        cell.fill = PatternFill("solid", fgColor=GRIS_CLARO)
                else:
                    for cell in row:
                        if zebra:
                            cell.fill = PatternFill("solid", fgColor="DBEAFE")
            ws2.freeze_panes = "A6"
            autofit(ws2)

            # ── Hoja 3: Clientes No Aptos ──
            df3 = pd.DataFrame(no_aptos).rename(columns={
                "cuenta_contrato":   "Cuenta Contrato",
                "instalacion":       "Instalación",
                "tarifa_referencia": "Tarifa Ref.",
                "observacion":       "Observación",
                "porcion":           "Porción",
                "unidad_predial":    "Unidad Predial",
                "meses_en_ventana":  "Meses Ventana",
                "estado_inicial":    "Estado Inicial",
            })
            df3.to_excel(writer, sheet_name="Clientes No Aptos", index=False, startrow=4)
            ws3 = writer.sheets["Clientes No Aptos"]
            nc3 = len(df3.columns)
            titulo_hoja(ws3, "CONTUGAS — Clientes No Aptos",
                        f"Clientes excluidos del proceso  |  Total: {len(df3):,}  |  {fecha_lima.strftime('%d/%m/%Y %H:%M')}", nc3)
            ws3.row_dimensions[5].height = 32
            estilo_header(ws3, 5, 1, nc3, "92400E", BLANCO)
            for i, row in enumerate(ws3.iter_rows(min_row=6, max_row=5 + len(df3), min_col=1, max_col=nc3)):
                for cell in row:
                    cell.fill      = PatternFill("solid", fgColor=AMBAR_CLARO if i % 2 == 0 else BLANCO)
                    cell.font      = Font(size=9, name="Calibri", color="374151")
                    cell.alignment = Alignment(vertical="center", wrap_text=True)
                    cell.border    = border_thin
            ws3.freeze_panes = "A6"
            autofit(ws3)

            # ── Hoja 4: Anomalías ──
            df4 = pd.DataFrame(anomalias).rename(columns={
                "cuenta_contrato": "Cuenta Contrato",
                "instalacion":     "Instalación",
                "fecha":           "Fecha",
                "tipo_anomalia":   "Tipo Anomalía",
            })
            df4.to_excel(writer, sheet_name="Anomalías", index=False, startrow=4)
            ws4 = writer.sheets["Anomalías"]
            nc4 = len(df4.columns)
            titulo_hoja(ws4, "CONTUGAS — Anomalías Detectadas",
                        f"Registros con comportamiento inusual  |  Total: {len(df4):,}  |  {fecha_lima.strftime('%d/%m/%Y %H:%M')}", nc4)
            ws4.row_dimensions[5].height = 32
            estilo_header(ws4, 5, 1, nc4, "7F1D1D", BLANCO)
            for i, row in enumerate(ws4.iter_rows(min_row=6, max_row=5 + len(df4), min_col=1, max_col=nc4)):
                for cell in row:
                    cell.fill      = PatternFill("solid", fgColor=ROJO_CLARO if i % 2 == 0 else BLANCO)
                    cell.font      = Font(size=9, name="Calibri", color="374151")
                    cell.alignment = Alignment(vertical="center")
                    cell.border    = border_thin
            ws4.freeze_panes = "A6"
            autofit(ws4)

            # ── Hoja 5: Matriz de Recategorización ──
            if matriz:
                df5 = pd.DataFrame(matriz).rename(columns={
                    "instalacion"        : "Instalación",
                    "cuenta_contrato"    : "Cuenta Contrato",
                    "porcion"            : "Porción",
                    "tipo_tarifa"        : "Tipo Tarifa",
                    "cf_mes_historico"   : "CF Hist.",
                    "fl_mes_historico"   : "FL Hist.",
                    "cl_mes_historico"   : "CL Hist.",
                    "cf_mes_1"           : "CF M1",
                    "fl_mes_1"           : "FL M1",
                    "dc_mes_1"           : "DC M1",
                    "cl_mes_1"           : "CL M1",
                    "tarifa_mes_1"       : "Tarifa M1",
                    "cf_mes_2"           : "CF M2",
                    "fl_mes_2"           : "FL M2",
                    "dc_mes_2"           : "DC M2",
                    "cl_mes_2"           : "CL M2",
                    "tarifa_mes_2"       : "Tarifa M2",
                    "cf_mes_3"           : "CF M3",
                    "fl_mes_3"           : "FL M3",
                    "dc_mes_3"           : "DC M3",
                    "cl_mes_3"           : "CL M3",
                    "tarifa_mes_3"       : "Tarifa M3",
                    "cf_mes_4"           : "CF M4",
                    "fl_mes_4"           : "FL M4",
                    "dc_mes_4"           : "DC M4",
                    "cl_mes_4"           : "CL M4",
                    "tarifa_mes_4"       : "Tarifa M4",
                    "cf_mes_5"           : "CF M5",
                    "fl_mes_5"           : "FL M5",
                    "dc_mes_5"           : "DC M5",
                    "cl_mes_5"           : "CL M5",
                    "tarifa_mes_5"       : "Tarifa M5",
                    "cf_mes_6"           : "CF M6",
                    "fl_mes_6"           : "FL M6",
                    "dc_mes_6"           : "DC M6",
                    "cl_mes_6"           : "CL M6",
                    "lectura_anterior"   : "Lec. Anterior",
                    "lectura_actual"     : "Lec. Actual",
                    "factor_correccion"  : "Factor Corrección",
                    "total_consumo"      : "Total Consumo",
                    "total_dias"         : "Total Días",
                    "promedio_diario"    : "Prom. Diario",
                    "promedio_mensual"   : "Prom. Mensual",
                    "promedio_redondeado": "Prom. Redondeado",
                    "tarifa_actual"      : "Tarifa Actual",
                    "tarifa_nueva"       : "Tarifa Nueva",
                    "recategorizar"      : "¿Recategorizar?",
                    "rango_consumo"      : "Rango Consumo",
                    "comportamiento"     : "Comportamiento",
                })
                df5.to_excel(writer, sheet_name="Matriz Recategorización", index=False, startrow=4)
                ws5 = writer.sheets["Matriz Recategorización"]
                nc5 = len(df5.columns)
                titulo_hoja(ws5, "CONTUGAS — Matriz de Recategorización",
                            f"Detalle por instalación · {len(df5):,} registros · {fecha_lima.strftime('%d/%m/%Y %H:%M')}", nc5)
                ws5.row_dimensions[5].height = 32
                estilo_header(ws5, 5, 1, nc5, "1E3A5F", BLANCO)
                col_rec = df5.columns.get_loc("¿Recategorizar?") + 1
                for i, row in enumerate(ws5.iter_rows(min_row=6, max_row=5 + len(df5), min_col=1, max_col=nc5)):
                    zebra = i % 2 == 0
                    for cell in row:
                        cell.fill      = PatternFill("solid", fgColor=AZUL_ROW if zebra else BLANCO)
                        cell.font      = Font(size=9, name="Calibri", color="374151")
                        cell.alignment = Alignment(vertical="center")
                        cell.border    = border_thin
                        if cell.column == col_rec:
                            if cell.value == "Sí":
                                cell.fill = PatternFill("solid", fgColor=VERDE_CLARO)
                                cell.font = Font(size=9, name="Calibri", color=VERDE, bold=True)
                            elif cell.value == "No":
                                cell.fill = PatternFill("solid", fgColor=GRIS_CLARO)
                                cell.font = Font(size=9, name="Calibri", color="6B7280")
                ws5.freeze_panes = "A6"
                autofit(ws5)

        output.seek(0)
        fecha_str = fecha_lima.strftime("%Y%m%d_%H%M")
        response = HttpResponse(
            output.read(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        response["Content-Disposition"] = f'attachment; filename="recategorizacion_{fecha_str}.xlsx"'
        return response


class ComparativaView(APIView):
    permission_classes = [IsAuthenticated]

    def get(self, request):
        id1 = request.query_params.get("id1")
        id2 = request.query_params.get("id2")

        if not id1 or not id2:
            return Response({"error": "Se requieren id1 e id2"}, status=400)

        try:
            if request.user.rol == 'admin':
                imp1 = ResultadoImportacion.objects.get(id=id1)
                imp2 = ResultadoImportacion.objects.get(id=id2)
            else:
                imp1 = ResultadoImportacion.objects.get(id=id1, usuario=request.user)
                imp2 = ResultadoImportacion.objects.get(id=id2, usuario=request.user)
        except ResultadoImportacion.DoesNotExist:
            return Response({"error": "Importación no encontrada"}, status=404)

        def resumen(imp):
            return {
                "id":              imp.id,
                "fecha":           imp.fecha_importacion.astimezone(
                    __import__('zoneinfo').ZoneInfo('America/Lima')
                ).strftime("%d/%m/%Y %H:%M"),
                "total_registros": imp.total_registros,
                "recategorizados": imp.recategorizados,
                "sin_cambios":     imp.sin_cambios,
                "no_aptos":        imp.no_aptos,
                "anomalias":       imp.anomalias,
                "distribucion":    list(
                    Cliente.objects.filter(importacion=imp)
                    .values("tarifa_nueva")
                    .annotate(cantidad=Count("id"))
                ),
            }

        return Response({"periodo1": resumen(imp1), "periodo2": resumen(imp2)})


class FiltrosDashboardView(APIView):
    permission_classes = [IsAuthenticated]

    def get(self, request):
        porcion        = request.query_params.get('porcion', '')
        fecha_desde    = request.query_params.get('fecha_desde', '')
        fecha_hasta    = request.query_params.get('fecha_hasta', '')
        importacion_id = request.query_params.get('importacion_id', '')

        if importacion_id:
            try:
                imp = ResultadoImportacion.objects.get(id=importacion_id, usuario=request.user)
            except ResultadoImportacion.DoesNotExist:
                return Response({"error": "Importación no encontrada"}, status=404)
        else:
            imp = ResultadoImportacion.objects.filter(usuario=request.user).order_by('-fecha_importacion').first()
            if not imp:
                return Response({"porciones": [], "periodos": [], "stats": None})

        porciones_disponibles = list(
            ConsumoMensual.objects.filter(importacion=imp)
            .exclude(porcion__isnull=True).exclude(porcion='')
            .values_list('porcion', flat=True)
            .distinct().order_by('porcion')
        )
        periodos_disponibles = list(
            ConsumoMensual.objects.filter(importacion=imp)
            .values_list('periodo', flat=True)
            .distinct().order_by('periodo')
        )

        if not porcion and not fecha_desde and not fecha_hasta:
            return Response({
                "porciones": porciones_disponibles,
                "periodos":  periodos_disponibles,
                "stats":     None,
            })

        qs = ConsumoMensual.objects.filter(importacion=imp)
        if porcion:
            qs = qs.filter(porcion=porcion)
        if fecha_desde:
            qs = qs.filter(periodo__gte=fecha_desde)
        if fecha_hasta:
            qs = qs.filter(periodo__lte=fecha_hasta)

        cuentas_filtradas = qs.values_list('cuenta_contrato', flat=True).distinct()
        clientes_qs       = Cliente.objects.filter(importacion=imp, cuenta_contrato__in=cuentas_filtradas)
        total             = clientes_qs.count()
        recategorizados   = clientes_qs.filter(estado='Recategorizado').count()
        sin_cambios       = clientes_qs.filter(estado='Sin cambio').count()

        no_aptos_qs    = ClienteNoApto.objects.filter(importacion=imp)
        if porcion:
            no_aptos_qs = no_aptos_qs.filter(porcion=porcion)
        no_aptos_count = no_aptos_qs.count()

        distribucion = list(clientes_qs.values('tarifa_nueva').annotate(cantidad=Count('id')))
        cambios      = list(clientes_qs.values('tarifa_anterior', 'tarifa_nueva').annotate(cantidad_clientes=Count('id')))
        if total > 0:
            for c in cambios:
                c['porcentaje'] = round(c['cantidad_clientes'] / total * 100, 2)

        consumo_por_periodo = list(
            qs.values('periodo')
            .annotate(consumo_total=Sum('consumo'), dias_total=Sum('dias'))
            .order_by('periodo')
        )

        return Response({
            "porciones": porciones_disponibles,
            "periodos":  periodos_disponibles,
            "stats": {
                "total_clientes":          total,
                "recategorizados":         recategorizados,
                "sin_cambios":             sin_cambios,
                "no_aptos":                no_aptos_count,
                "distribucion_categorias": distribucion,
                "cambios_tarifarios":      cambios,
                "consumo_por_periodo":     consumo_por_periodo,
            }
        })


class MatrizRecategorizacionView(APIView):
    permission_classes = [IsAuthenticated]

    def get(self, request):
        importacion_id = request.query_params.get('importacion_id', '')
        search         = request.query_params.get('search', '')
        recategorizar  = request.query_params.get('recategorizar', '')
        page           = int(request.query_params.get('page', 1))
        size           = int(request.query_params.get('page_size', 50))

        es_admin = request.user.rol == 'admin'
        if importacion_id:
            try:
                imp = ResultadoImportacion.objects.get(id=importacion_id) if es_admin else ResultadoImportacion.objects.get(id=importacion_id, usuario=request.user)
            except ResultadoImportacion.DoesNotExist:
                imp = ResultadoImportacion.objects.order_by('-fecha_importacion').first() if es_admin else ResultadoImportacion.objects.filter(usuario=request.user).order_by('-fecha_importacion').first()
        else:
            imp = ResultadoImportacion.objects.order_by('-fecha_importacion').first() if es_admin else ResultadoImportacion.objects.filter(usuario=request.user).order_by('-fecha_importacion').first()

        if not imp:
            return Response({"count": 0, "results": []})

        qs = MatrizRecategorizacion.objects.filter(importacion=imp)
        if search:
            qs = qs.filter(Q(instalacion__icontains=search) | Q(cuenta_contrato__icontains=search))
        if recategorizar:
            qs = qs.filter(recategorizar=recategorizar)

        total  = qs.count()
        offset = (page - 1) * size
        rows   = qs[offset:offset + size]

        results = []
        for r in rows:
            results.append({
                "instalacion"        : r.instalacion,
                "cuenta_contrato"    : r.cuenta_contrato,
                "porcion"            : r.porcion or "-",
                "tipo_tarifa"        : r.tipo_tarifa,
                "cf_mes_historico"   : r.cf_mes_historico,
                "fl_mes_historico"   : str(r.fl_mes_historico) if r.fl_mes_historico else None,
                "cl_mes_historico"   : r.cl_mes_historico,
                "cf_mes_1"           : r.cf_mes_1,  "fl_mes_1": str(r.fl_mes_1) if r.fl_mes_1 else None,  "dc_mes_1": r.dc_mes_1,  "cl_mes_1": r.cl_mes_1,  "tarifa_mes_1": r.tarifa_mes_1,
                "cf_mes_2"           : r.cf_mes_2,  "fl_mes_2": str(r.fl_mes_2) if r.fl_mes_2 else None,  "dc_mes_2": r.dc_mes_2,  "cl_mes_2": r.cl_mes_2,  "tarifa_mes_2": r.tarifa_mes_2,
                "cf_mes_3"           : r.cf_mes_3,  "fl_mes_3": str(r.fl_mes_3) if r.fl_mes_3 else None,  "dc_mes_3": r.dc_mes_3,  "cl_mes_3": r.cl_mes_3,  "tarifa_mes_3": r.tarifa_mes_3,
                "cf_mes_4"           : r.cf_mes_4,  "fl_mes_4": str(r.fl_mes_4) if r.fl_mes_4 else None,  "dc_mes_4": r.dc_mes_4,  "cl_mes_4": r.cl_mes_4,  "tarifa_mes_4": r.tarifa_mes_4,
                "cf_mes_5"           : r.cf_mes_5,  "fl_mes_5": str(r.fl_mes_5) if r.fl_mes_5 else None,  "dc_mes_5": r.dc_mes_5,  "cl_mes_5": r.cl_mes_5,  "tarifa_mes_5": r.tarifa_mes_5,
                "cf_mes_6"           : r.cf_mes_6,  "fl_mes_6": str(r.fl_mes_6) if r.fl_mes_6 else None,  "dc_mes_6": r.dc_mes_6,  "cl_mes_6": r.cl_mes_6,
                "lectura_anterior"   : r.lectura_anterior,
                "lectura_actual"     : r.lectura_actual,
                "factor_correccion"  : r.factor_correccion,
                "total_consumo"      : round(r.total_consumo, 2),
                "total_dias"         : r.total_dias,
                "promedio_diario"    : round(r.promedio_diario, 4),
                "promedio_mensual"   : round(r.promedio_mensual, 2),
                "promedio_redondeado": r.promedio_redondeado,
                "tarifa_actual"      : r.tarifa_actual,
                "tarifa_nueva"       : r.tarifa_nueva,
                "recategorizar"      : r.recategorizar,
                "rango_consumo"      : r.rango_consumo,
                "comportamiento"     : r.comportamiento,
            })

        return Response({"count": total, "results": results})


class ClienteNoAptoListView(APIView):
    permission_classes = [IsAuthenticated]

    def get(self, request):
        importacion_id = request.query_params.get('importacion_id', '')
        search         = request.query_params.get('search', '')
        page           = int(request.query_params.get('page', 1))
        size           = int(request.query_params.get('page_size', 50))

        es_admin = request.user.rol == 'admin'
        if importacion_id:
            try:
                imp = ResultadoImportacion.objects.get(id=importacion_id) if es_admin else ResultadoImportacion.objects.get(id=importacion_id, usuario=request.user)
            except ResultadoImportacion.DoesNotExist:
                imp = ResultadoImportacion.objects.order_by('-fecha_importacion').first() if es_admin else ResultadoImportacion.objects.filter(usuario=request.user).order_by('-fecha_importacion').first()
        else:
            imp = ResultadoImportacion.objects.order_by('-fecha_importacion').first() if es_admin else ResultadoImportacion.objects.filter(usuario=request.user).order_by('-fecha_importacion').first()

        if not imp:
            return Response({"count": 0, "results": []})

        qs = ClienteNoApto.objects.filter(importacion=imp)
        if search:
            qs = qs.filter(Q(instalacion__icontains=search) | Q(cuenta_contrato__icontains=search))

        total  = qs.count()
        offset = (page - 1) * size
        rows   = qs[offset:offset + size]

        results = [{
            "instalacion"      : r.instalacion,
            "cuenta_contrato"  : r.cuenta_contrato,
            "tarifa_referencia": r.tarifa_referencia,
            "observacion"      : r.observacion,
            "porcion"          : r.porcion or "—",
            "meses_en_ventana" : r.meses_en_ventana,
            "estado_inicial"   : r.estado_inicial or "No apto",
        } for r in rows]

        return Response({"count": total, "results": results})